from __future__ import annotations

from pathlib import Path

import polars as pl
from openpyxl import load_workbook

from carta.excel import write_xlsx


def test_writes_header_data_and_freezes(tmp_path: Path, sample_df: pl.DataFrame) -> None:
    out = tmp_path / "out.xlsx"
    write_xlsx(sample_df, out, sheet_name="Data")
    assert out.is_file()

    wb = load_workbook(out)
    assert wb.sheetnames == ["Data"]
    ws = wb["Data"]
    assert ws.freeze_panes == "A2"
    # Header row
    assert [c.value for c in ws[1]] == ["student_id", "program_code", "credits"]
    assert ws["A1"].font.bold is True
    # First data row
    assert [c.value for c in ws[2]] == [1001, "CS", 15]
    # Three data rows
    assert ws.max_row == 4


def test_empty_df_writes_just_an_empty_sheet(tmp_path: Path) -> None:
    out = tmp_path / "empty.xlsx"
    write_xlsx(pl.DataFrame(), out, sheet_name="Data")
    wb = load_workbook(out)
    ws = wb["Data"]
    assert ws.max_row == 1  # empty sheet has one row by default
    assert ws["A1"].value is None


def test_column_widths_set(tmp_path: Path) -> None:
    df = pl.DataFrame({"short": ["a"], "longer_header_name": ["x"]})
    out = tmp_path / "w.xlsx"
    write_xlsx(df, out, sheet_name="Data")
    wb = load_workbook(out)
    ws = wb["Data"]
    a = ws.column_dimensions["A"].width
    b = ws.column_dimensions["B"].width
    assert a is not None and a > 0
    assert b is not None and b > a  # longer header => wider column


def test_none_cells_written_as_empty_string(tmp_path: Path) -> None:
    df = pl.DataFrame({"x": [None, "a"]}, schema={"x": pl.Utf8})
    out = tmp_path / "n.xlsx"
    write_xlsx(df, out, sheet_name="Data")
    wb = load_workbook(out)
    ws = wb["Data"]
    # openpyxl reads empty string as None on read-back; the key thing is no crash
    assert ws["A2"].value in (None, "")
    assert ws["A3"].value == "a"
