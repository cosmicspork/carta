"""End-to-end CLI tests using subprocess.

Tests that need a working connection use a sqlite DATABASE_URL — no driver
stub or PYTHONPATH shadowing needed.
"""

from __future__ import annotations

import os
import sqlite3
import subprocess
import sys
from pathlib import Path

from tests.conftest import write_project

REPO = Path(__file__).resolve().parent.parent


def _setup_sqlite(db_path: Path) -> None:
    """Create a small `students` table the test SQL queries against."""
    c = sqlite3.connect(db_path)
    c.execute("CREATE TABLE students (id INTEGER, name TEXT)")
    c.executemany("INSERT INTO students VALUES (?, ?)", [(1, "alpha"), (2, "beta")])
    c.commit()
    c.close()


def _project(
    tmp_path: Path, *, active: list[str], with_sqlite: bool = False
) -> Path:
    """Build a project tree. If `with_sqlite`, also create the test DB and point
    DATABASE_URL at it."""
    if with_sqlite:
        sqlite_path = tmp_path / "carta_test.sqlite"
        url = f"sqlite:///{sqlite_path}"
    else:
        url = "sqlite:///:memory:"  # never actually connects in no-active tests
    proj = write_project(tmp_path / "proj", active=active, database_url=url)
    if with_sqlite:
        _setup_sqlite(tmp_path / "carta_test.sqlite")
    return proj


def _run(cwd: Path, *args: str) -> subprocess.CompletedProcess[str]:
    env = os.environ.copy()
    env["PYTHONPATH"] = os.pathsep.join([str(REPO), env.get("PYTHONPATH", "")]).rstrip(os.pathsep)
    return subprocess.run(
        [sys.executable, "-m", "carta", *args],
        cwd=cwd,
        env=env,
        capture_output=True,
        text=True,
    )


_GOOD_PY = (
    "definition = {\n"
    "    'key': 'enrollment',\n"
    "    'recipients': ['a@x.com'],\n"
    "    'subject': 'Enrollment {run_date}',\n"
    "    'body': 'Rows: {row_count}',\n"
    "}\n"
)
_GOOD_SQL = "SELECT id, name FROM students ORDER BY id\n"


def test_no_active_keys_exit_0(tmp_path: Path) -> None:
    proj = _project(tmp_path, active=[])
    r = _run(proj)
    assert r.returncode == 0, r.stdout + r.stderr
    assert "Nothing to run." in r.stdout


def test_ghost_active_key_exit_1(tmp_path: Path) -> None:
    proj = _project(tmp_path, active=["ghost"])
    r = _run(proj)
    assert r.returncode == 1, r.stdout
    assert "ghost" in r.stdout
    assert "ACTIVE" in r.stdout


def test_missing_sql_aborts_exit_2(tmp_path: Path) -> None:
    proj = _project(tmp_path, active=["enrollment"])
    (proj / "definitions" / "enrollment.py").write_text(_GOOD_PY)
    # No sibling .sql
    r = _run(proj)
    assert r.returncode == 2, r.stdout
    assert "sibling SQL file not found" in r.stdout
    assert "Aborting" in r.stdout


def test_broken_definition_isolated_exit_1(tmp_path: Path) -> None:
    proj = _project(tmp_path, active=["enrollment"], with_sqlite=True)
    (proj / "definitions" / "broken.py").write_text("raise RuntimeError('boom')\n")
    (proj / "definitions" / "enrollment.py").write_text(_GOOD_PY)
    (proj / "definitions" / "enrollment.sql").write_text(_GOOD_SQL)
    r = _run(proj)
    assert r.returncode == 1, r.stdout
    assert "broken.py" in r.stdout
    assert "ok    [enrollment]" in r.stdout


def test_duplicate_keys_abort_exit_2(tmp_path: Path) -> None:
    proj = _project(tmp_path, active=["enrollment"])
    (proj / "definitions" / "a.py").write_text(_GOOD_PY)
    (proj / "definitions" / "a.sql").write_text(_GOOD_SQL)
    (proj / "definitions" / "b.py").write_text(_GOOD_PY)
    (proj / "definitions" / "b.sql").write_text(_GOOD_SQL)
    r = _run(proj)
    assert r.returncode == 2, r.stdout
    assert "duplicate key" in r.stdout


def test_duplicate_filename_aborts_exit_2(tmp_path: Path) -> None:
    proj = _project(tmp_path, active=["one", "two"])
    base = _GOOD_PY.replace("'enrollment'", "'one'")
    (proj / "definitions" / "one.py").write_text(
        base.replace("    'body'", "    'filename': 'shared.xlsx',\n    'body'")
    )
    (proj / "definitions" / "one.sql").write_text(_GOOD_SQL)
    (proj / "definitions" / "two.py").write_text(
        _GOOD_PY.replace("'enrollment'", "'two'").replace(
            "    'body'", "    'filename': 'shared.xlsx',\n    'body'"
        )
    )
    (proj / "definitions" / "two.sql").write_text(_GOOD_SQL)
    r = _run(proj)
    assert r.returncode == 2, r.stdout
    assert "also produced by" in r.stdout


def test_clean_happy_path_with_sqlite_backend(tmp_path: Path) -> None:
    proj = _project(tmp_path, active=["enrollment"], with_sqlite=True)
    (proj / "definitions" / "enrollment.py").write_text(_GOOD_PY)
    (proj / "definitions" / "enrollment.sql").write_text(_GOOD_SQL)
    r = _run(proj)
    assert r.returncode == 0, r.stdout + r.stderr
    assert "ok    [enrollment] 2 rows" in r.stdout

    run_dirs = list((proj / "out").iterdir())
    assert len(run_dirs) == 1
    files = sorted(p.name for p in run_dirs[0].iterdir())
    assert "enrollment.xlsx" in files
    assert "enrollment.eml" in files

    from openpyxl import load_workbook

    wb = load_workbook(run_dirs[0] / "enrollment.xlsx")
    ws = wb["Data"]
    assert [c.value for c in ws[1]] == ["id", "name"]
    assert [c.value for c in ws[2]] == [1, "alpha"]
    assert [c.value for c in ws[3]] == [2, "beta"]


def test_only_filter_runs_subset(tmp_path: Path) -> None:
    proj = _project(tmp_path, active=["one", "two"], with_sqlite=True)
    (proj / "definitions" / "one.py").write_text(_GOOD_PY.replace("'enrollment'", "'one'"))
    (proj / "definitions" / "one.sql").write_text(_GOOD_SQL)
    (proj / "definitions" / "two.py").write_text(_GOOD_PY.replace("'enrollment'", "'two'"))
    (proj / "definitions" / "two.sql").write_text(_GOOD_SQL)
    r = _run(proj, "--only", "one")
    assert r.returncode == 0, r.stdout + r.stderr
    assert "ok    [one]" in r.stdout
    assert "[two]" not in r.stdout


def test_only_unknown_key_warned(tmp_path: Path) -> None:
    proj = _project(tmp_path, active=["enrollment"], with_sqlite=True)
    (proj / "definitions" / "enrollment.py").write_text(_GOOD_PY)
    (proj / "definitions" / "enrollment.sql").write_text(_GOOD_SQL)
    r = _run(proj, "--only", "ghost")
    assert "'ghost' is not in ACTIVE" in r.stdout
    assert r.returncode == 0
    assert "Nothing to run." in r.stdout


def test_param_binding_through_full_pipeline(tmp_path: Path) -> None:
    """Parameter binding survives the round-trip through the real DB-API driver."""
    proj = _project(tmp_path, active=["bound"], with_sqlite=True)
    (proj / "definitions" / "bound.py").write_text(
        "definition = {\n"
        "    'key': 'bound',\n"
        "    'recipients': ['a@x.com'],\n"
        "    'subject': 'S',\n"
        "    'body': 'B',\n"
        "    'params': {'wanted_id': 2},\n"
        "}\n"
    )
    (proj / "definitions" / "bound.sql").write_text(
        "SELECT id, name FROM students WHERE id = :wanted_id\n"
    )
    r = _run(proj)
    assert r.returncode == 0, r.stdout + r.stderr
    assert "ok    [bound] 1 rows" in r.stdout

    from openpyxl import load_workbook

    run_dir = next((proj / "out").iterdir())
    ws = load_workbook(run_dir / "bound.xlsx")["Data"]
    assert [c.value for c in ws[2]] == [2, "beta"]
    assert ws.max_row == 2


def test_validation_error_isolated_exit_1(tmp_path: Path) -> None:
    proj = _project(tmp_path, active=["enrollment", "broken_contract"], with_sqlite=True)
    (proj / "definitions" / "enrollment.py").write_text(_GOOD_PY)
    (proj / "definitions" / "enrollment.sql").write_text(_GOOD_SQL)
    (proj / "definitions" / "broken_contract.py").write_text(
        "definition = {'key': 'broken_contract', 'recipients': [], "
        "'subject': 'S', 'body': 'B'}\n"
    )
    (proj / "definitions" / "broken_contract.sql").write_text(_GOOD_SQL)
    r = _run(proj)
    assert r.returncode == 1, r.stdout
    assert "broken_contract" in r.stdout
    assert "ok    [enrollment]" in r.stdout


def test_invalid_database_url_aborts_exit_2(tmp_path: Path) -> None:
    proj = _project(tmp_path, active=["enrollment"])
    # Hand-write a config with a bad URL so write_project's default doesn't apply
    (proj / "config.py").write_text(
        "from pathlib import Path\n"
        "REPO_ROOT = Path(__file__).resolve().parent\n"
        "DEFINITIONS_DIR = REPO_ROOT / 'definitions'\n"
        "OUTPUT_DIR = REPO_ROOT / 'out'\n"
        "DATABASE_URL = 'not a valid url'\n"
        "KEYRING_SERVICE = 'carta_test'\n"
        "ACTIVE = ['enrollment']\n"
    )
    (proj / "definitions" / "enrollment.py").write_text(_GOOD_PY)
    (proj / "definitions" / "enrollment.sql").write_text(_GOOD_SQL)
    r = _run(proj)
    assert r.returncode == 2, r.stdout
    assert "could not connect to database" in r.stdout
